"""
CipherVault — Professional Presentation Generator (Theme 1: Linear Dark Vault)
Run: python generate_ppt.py
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
import os

# ==============================================================================
# LOCKED COURSE, GUIDE & STUDENT DETAILS (PROFESSIONAL MATRIX)
# ==============================================================================
COURSE_NAME  = "Full Stack Web Development Project using ASP.NET Core"
COURSE_CODE  = "DSC-M-BCA-471P"
STUDENT_NAME = "Kaswan Yogesh Subhashchandra"
ROLL_NO      = "32"
COLLEGE      = "Lokmanya College of Computer Application"
GUIDE_NAME   = "Prof. Sejal Shah"
DATE         = "06-10-2026"
LIVE_URL     = "https://ciphervault.tryasp.net"
GITHUB_URL   = "https://github.com/yogeshkaswan377-dev/CipherVault"
SCREENSHOTS_DIR = "docs/screenshots"
OUTPUT_FILE     = "CipherVault_Professional_Presentation.pptx"

# ==============================================================================
# THEME 1 TOKENS: "Linear Dark Vault" (App-Aligned & Professional)
# ==============================================================================
BG_OBSIDIAN  = RGBColor(11, 15, 25)    # #0B0F19 Deep Obsidian Navy
CARD_BG      = RGBColor(26, 35, 51)    # #1A2333 Dark Glass Card
CARD_ALT     = RGBColor(30, 41, 59)    # #1E293B Slate Card
INSIGHT_BG   = RGBColor(23, 37, 84)    # #172554 Deep Royal Insight Box
BORDER_SLATE = RGBColor(51, 65, 85)    # #334155 Subtle Border
ACCENT_INDIGO= RGBColor(94, 106, 210)  # #5E6AD2 Linear Indigo
ACCENT_VIOLET= RGBColor(139, 149, 246) # #8B95F6 Secondary Glow
SUCCESS      = RGBColor(16, 185, 129)  # #10B981 Emerald Green (Encrypted/Pass)
WARNING      = RGBColor(245, 158, 11)  # #F59E0B Amber (Blank-preserve rule)
DANGER       = RGBColor(239, 68, 68)   # #EF4444 Rose (IDOR / 401 / 404)
CYAN_GLOW    = RGBColor(6, 182, 212)   # #06B6D4 Tech Cyan (API & JWT)
TEXT_WHITE   = RGBColor(248, 250, 252) # #F8FAFC Crisp White
TEXT_MUTED   = RGBColor(148, 163, 184) # #94A3B8 Muted Silver
TEXT_LO      = RGBColor(107, 114, 128) # #6B7280 Footers

FONT_HEADING = "Segoe UI"
FONT_BODY    = "Segoe UI"
FONT_MONO    = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)

prs = Presentation()
prs.slide_width = SW
prs.slide_height = SH
BLANK = prs.slide_layouts[6]
slide_num = 0

# ==============================================================================
# PROFESSIONAL HELPER FUNCTIONS
# ==============================================================================
def new_slide(bg=BG_OBSIDIAN):
    global slide_num
    slide_num += 1
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = bg
    try:
        transition_xml = parse_xml(
            '<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med" advClick="1">'
            '  <p:extLst>'
            '    <p:ext uri="{D9489498-53F2-484D-A90E-23E2B614D64A}">'
            '      <p159:morph xmlns:p159="http://schemas.microsoft.com/office/powerpoint/2015/09/main" option="byObject"/>'
            '    </p:ext>'
            '  </p:extLst>'
            '</p:transition>'
        )
        s._element.insert(1, transition_xml)
    except Exception:
        pass
    return s

def add_rect(slide, left, top, width, height, fill=None, line=None, line_w=1, rounded=True, radius=0.06):
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(shp_type, Inches(left), Inches(top), Inches(width), Inches(height))
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    shp.shadow.inherit = False
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    return shp

def add_card(slide, left, top, width, height, border_color=BORDER_SLATE, fill_color=CARD_BG, name=None):
    shape = add_rect(slide, left, top, width, height, fill=fill_color, line=border_color, line_w=1.75, rounded=True, radius=0.04)
    if name:
        shape.name = name
    return shape

def add_text(slide, text, left, top, width, height, size=14, color=TEXT_WHITE, bold=False, align=PP_ALIGN.LEFT, name=None, font=FONT_BODY):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    if name:
        tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = font
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    return tb

def add_bullets(slide, items, left, top, width, height, size=12, color=TEXT_WHITE, spacing=6):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(spacing)
        if isinstance(item, tuple):
            prefix, rest = item[0], item[1]
            p_color = item[2] if len(item) > 2 else SUCCESS
            r1 = p.add_run()
            r1.text = prefix
            r1.font.name = FONT_BODY
            r1.font.size = Pt(size)
            r1.font.bold = True
            r1.font.color.rgb = p_color

            r2 = p.add_run()
            r2.text = rest
            r2.font.name = FONT_BODY
            r2.font.size = Pt(size)
            r2.font.bold = False
            r2.font.color.rgb = color
        else:
            p.text = item
            p.font.name = FONT_BODY
            p.font.size = Pt(size)
            p.font.color.rgb = color
    return tb

def feature_card(slide, left, top, width, height, icon, title, lines, accent=ACCENT_INDIGO):
    card = add_rect(slide, left, top, width, height, fill=CARD_BG, rounded=True, radius=0.08, line=BORDER_SLATE, line_w=1)
    add_rect(slide, left + 0.15, top, width - 0.3, 0.04, fill=accent)
    
    # Icon chip
    chip = add_rect(slide, left + 0.25, top + 0.22, 0.5, 0.5, fill=accent, line=None, rounded=True, radius=0.22)
    chip.text_frame.text = icon
    p = chip.text_frame.paragraphs[0]
    p.font.size = Pt(16)
    p.font.color.rgb = BG_OBSIDIAN
    p.alignment = PP_ALIGN.CENTER
    chip.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    add_text(slide, title, left + 0.85, top + 0.28, width - 1.0, 0.4, size=14, bold=True, color=TEXT_WHITE, font=FONT_HEADING)

    tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.9), Inches(width - 0.5), Inches(height - 1.0))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = 0
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = "· " + line
        p.font.size = Pt(10.5)
        p.font.color.rgb = TEXT_MUTED
        p.font.name = FONT_BODY
        p.space_after = Pt(4)

def add_header(slide, phase_badge, title_text, objective_text, badge_color=SUCCESS):
    add_card(slide, 0.55, 0.42, 2.75, 0.38, border_color=badge_color, fill_color=CARD_ALT, name="!!phase_badge")
    add_text(slide, phase_badge, 0.6, 0.45, 2.65, 0.32, size=10.5, color=badge_color, bold=True, align=PP_ALIGN.CENTER)

    add_text(slide, title_text, 3.42, 0.38, 8.5, 0.42, size=19, color=TEXT_WHITE, bold=True, name="!!slide_title")
    add_text(slide, objective_text, 3.42, 0.78, 8.5, 0.32, size=11, color=CYAN_GLOW, bold=True)

def add_insight_box(slide, label_title, insight_text, top=6.05, height=0.88):
    add_card(slide, 0.55, top, 12.233, height, border_color=CYAN_GLOW, fill_color=INSIGHT_BG, name="!!insight_strip")
    add_bullets(slide, [
        (f"💡 {label_title}:  ", insight_text, ACCENT_VIOLET)
    ], 0.75, top + 0.08, 11.85, height - 0.12, size=11.2, color=TEXT_WHITE, spacing=2)

def add_footer(slide, slide_num, total_slides=15):
    line = add_rect(slide, 0.55, 7.04, 12.233, 0.02, fill=BORDER_SLATE, line=None)

    footer_left = (
        f"{COURSE_NAME} ({COURSE_CODE})  |  Student: {STUDENT_NAME} (Roll: {ROLL_NO})  |  "
        f"Guide: {GUIDE_NAME}"
    )
    add_text(slide, footer_left, 0.55, 7.09, 10.5, 0.28, size=9, color=TEXT_LO)
    add_text(slide, f"Slide {slide_num} / {total_slides}", 11.1, 7.09, 1.78, 0.28, size=9.5, color=ACCENT_VIOLET, bold=True, align=PP_ALIGN.RIGHT)

def add_plot_frame(slide, filename, caption, left, top, width, height, border_color=CYAN_GLOW):
    add_card(slide, left, top, width, height, border_color=border_color, fill_color=CARD_BG)
    img_path = os.path.join(SCREENSHOTS_DIR, filename)
    pad = 0.1
    img_w = width - (2 * pad)
    img_h = height - 0.42

    if os.path.exists(img_path):
        slide.shapes.add_picture(img_path, Inches(left + pad), Inches(top + pad), width=Inches(img_w), height=Inches(img_h))
    else:
        add_card(slide, left + pad, top + pad, img_w, img_h, border_color=WARNING, fill_color=BG_OBSIDIAN)
        add_text(slide, f"[ Screenshot Missing: {filename} ]", left + pad + 0.15, top + (img_h / 2) - 0.2, img_w - 0.3, 0.5, size=10, color=WARNING, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, caption, left + 0.12, top + height - 0.31, width - 0.24, 0.26, size=9.5, color=TEXT_MUTED, bold=True, align=PP_ALIGN.CENTER)

def add_notes(slide, notes_text):
    slide.notes_slide.notes_text_frame.text = notes_text

def style_table(table, col_widths, font_size=10):
    for idx, w in enumerate(col_widths):
        table.columns[idx].width = Inches(w)
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if r_idx == 0:
                cell.fill.fore_color.rgb = CARD_ALT
            elif r_idx % 2 == 1:
                cell.fill.fore_color.rgb = CARD_BG
            else:
                cell.fill.fore_color.rgb = BG_OBSIDIAN
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT_BODY
                p.font.size = Pt(font_size if r_idx > 0 else font_size + 0.5)
                p.font.bold = (r_idx == 0)
                p.font.color.rgb = ACCENT_VIOLET if r_idx == 0 else TEXT_WHITE
                p.alignment = PP_ALIGN.LEFT if c_idx <= 1 else PP_ALIGN.CENTER


# ==============================================================================
# BUILD 15-SLIDE PROFESSIONAL CIPHERVAULT PRESENTATION
# ==============================================================================
def build_presentation():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    blank_layout = prs.slide_layouts[6]

    # --------------------------------------------------------------------------
    # SLIDE 1: HERO TITLE SLIDE
    # --------------------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    new_slide() # triggers slide count
    s1.background.fill.solid(); s1.background.fill.fore_color.rgb = BG_OBSIDIAN

    add_card(s1, 0.55, 0.55, 12.233, 6.2, border_color=ACCENT_INDIGO, fill_color=CARD_BG)
    add_card(s1, 0.9, 0.88, 7.5, 0.42, border_color=ACCENT_VIOLET, fill_color=CARD_ALT)
    add_text(s1, f"{COURSE_NAME.upper()}  •  CODE: {COURSE_CODE}", 0.95, 0.93, 7.4, 0.32, size=10.5, color=ACCENT_VIOLET, bold=True, align=PP_ALIGN.CENTER)

    add_text(s1, "CipherVault", 0.9, 1.45, 11.5, 0.8, size=44, color=TEXT_WHITE, bold=True)
    add_text(s1, "Personal Encrypted Password & Secure Notes Manager", 0.9, 2.22, 11.5, 0.4, size=16, color=SUCCESS, bold=True)
    add_text(
        s1,
        "A Professional Full Stack Web Application Built on ASP.NET Core 10 MVC & Entity Framework Core 10, "
        "Featuring AES-Authenticated Encryption at Rest, Ownership-Scoped IDOR Protection, and JWT-Secured REST API.",
        0.9, 2.75, 11.2, 0.7, size=11.8, color=TEXT_MUTED
    )

    add_card(s1, 0.9, 3.65, 11.533, 2.05, border_color=BORDER_SLATE, fill_color=CARD_ALT)
    add_bullets(s1, [
        ("Student Name & Roll No: ", f"{STUDENT_NAME}    |    Roll Number: {ROLL_NO}", SUCCESS),
        ("Course Title & Code: ", f"{COURSE_NAME}    |    {COURSE_CODE}", ACCENT_VIOLET),
        ("Project Guide: ", f"{GUIDE_NAME}    |    Academic Year: 2026", SUCCESS),
        ("Tech Stack: ", "ASP.NET Core 10 MVC  •  EF Core 10  •  SQL Server  •  ASP.NET Core Identity  •  Data Protection API (v10.0.7+)  •  Bootstrap 5", CYAN_GLOW),
    ], 1.1, 3.8, 11.1, 1.8, size=11.5, spacing=6)

    add_card(s1, 0.55, 6.85, 12.233, 0.45, border_color=BORDER_SLATE, fill_color=CARD_ALT)
    add_text(s1, "PROJECT ROADMAP:   [PHASE 1: MVC Setup & UI]   •   [PHASE 2: Forms & Validation]   •   [PHASE 3: Database & CRUD]   •   [PHASE 4: API & Deployment]", 0.7, 6.92, 11.9, 0.3, size=9.5, color=SUCCESS, bold=True, align=PP_ALIGN.CENTER)
    add_footer(s1, 1)
    add_notes(s1, f"Good morning Sir/Ma'am. My name is {STUDENT_NAME} (Roll No {ROLL_NO}). Under the guidance of {GUIDE_NAME}, my project is 'CipherVault', built under Full Stack Web Development using ASP.NET Core (Course Code: {COURSE_CODE}).")

    # --------------------------------------------------------------------------
    # SLIDE 2: 4-PHASE PROJECT ROADMAP & MATRIX
    # --------------------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    new_slide()
    s2.background.fill.solid(); s2.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s2, "PROJECT ROADMAP", "4-Phase Full Stack Development Matrix", f"Structured Execution Aligned with {COURSE_NAME} ({COURSE_CODE}) — Guide: {GUIDE_NAME}", badge_color=ACCENT_INDIGO)

    add_card(s2, 0.52, 1.2, 12.29, 4.72, border_color=ACCENT_INDIGO, fill_color=CARD_BG)
    t2_shape = s2.shapes.add_table(5, 4, Inches(0.55), Inches(1.25), Inches(12.233), Inches(4.62))
    t2 = t2_shape.table
    headers2 = ["Project Phase & Title", "Phase Objective", "Key Technical Implementation Tasks", "Completed Deliverables"]
    rows2 = [
        [
            "PHASE 1:\nMVC Setup and UI",
            "To configure the ASP.NET Core 10 architecture and build the responsive user interface shell.",
            "• Project setup targeting net10.0 with EF Core 10 & Identity\n• Program.cs with WebApplicationBuilder & Cookie policy (SecurePolicy=Always)\n• Base layout (_Layout.cshtml) with responsive Bootstrap 5 navbar\n• HomeController & placeholder [Authorize] DashboardController",
            "✔ Working Authentication\n✔ Base Layout & Navbar\n✔ Authorized Dashboard\n✔ Target net10.0 Confirmed"
        ],
        [
            "PHASE 2:\nForms and Validation",
            "To implement secure user interaction forms, data validation, and cryptographic utilities.",
            "• VaultItem model with categories (Password, Note, API Key, Credit Card, Other)\n• Separate Create/Update DTOs with validation annotations & custom ValidCategoryAttribute\n• Client-side password generator using window.crypto.getRandomValues()\n• Data Protection key persistence configuration (App_Data/keys)",
            "✔ Validated Create/Edit Forms\n✔ Crypto-Secure Generator\n✔ DTO Separation\n✔ Key Persistence Config"
        ],
        [
            "PHASE 3:\nDatabase and CRUD",
            "To build the database persistence layer, encrypted repository, and full CRUD lifecycle.",
            "• ApplicationDbContext extended with DbSet<VaultItem> & User relationship\n• Repository pattern implemented for VaultItem (IVaultItemRepository)\n• Service layer handling AES encryption via purpose strings (\"CipherVault.Secret/.Notes\")\n• IDOR protection: Ownership checks returning 404 (not 403) for non-owned items",
            "✔ Full Encrypted CRUD\n✔ Service & Repository Layers\n✔ IDOR 404 Enforcement\n✔ Search & Category Filter"
        ],
        [
            "PHASE 4:\nAPI and Deployment",
            "To expose a secure REST API with JWT Bearer auth and deploy to live production hosting.",
            "• REST API controllers under /api/v1/vault with [ApiController]\n• Dual authentication scheme separation: Cookies for MVC, JWT for API\n• Global ExceptionHandlingMiddleware returning JSON 500 / HTML error pages\n• Live deployment on MonsterASP (ciphervault.tryasp.net) via HTTPS",
            "✔ REST API + JWT Bearer\n✔ Global Error Middleware\n✔ Live HTTPS Deployment\n✔ Postman Collection"
        ]
    ]
    for c_i, h in enumerate(headers2):
        t2.cell(0, c_i).text = h
    for r_i, r_data in enumerate(rows2, start=1):
        for c_i, val in enumerate(r_data):
            t2.cell(r_i, c_i).text = val
    style_table(t2, [2.1, 2.5, 5.233, 2.4], font_size=9.8)
    for r_i in range(1, 5):
        t2.cell(r_i, 0).text_frame.paragraphs[0].font.color.rgb = SUCCESS
        t2.cell(r_i, 0).text_frame.paragraphs[0].font.bold = True
        t2.cell(r_i, 2).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        t2.cell(r_i, 3).text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT
        t2.cell(r_i, 3).text_frame.paragraphs[0].font.color.rgb = CYAN_GLOW

    add_insight_box(
        s2,
        "MODULAR FULL STACK ARCHITECTURE",
        "Each of the 4 phases is self-contained with strict separation between controllers, services, and repositories, ensuring clean maintainability and robust security adherence."
    )
    add_footer(s2, 2)
    add_notes(s2, "We followed a structured 4-phase roadmap: Phase 1 MVC setup, Phase 2 Forms & Validation, Phase 3 Database & CRUD, and Phase 4 API & Deployment.")

    # --------------------------------------------------------------------------
    # SLIDE 3: PHASE 1 — MVC SETUP, LAYOUT & AUTHENTICATION
    # --------------------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    new_slide()
    s3.background.fill.solid(); s3.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s3, "PHASE 1 • ARCHITECTURE", "Phase 1: MVC Setup, Base Layout & Authentication", "Objective: Configure ASP.NET Core 10 architecture and establish secure cookie authentication", badge_color=ACCENT_INDIGO)

    feature_card(s3, 0.7, 2.3, 3.8, 3.5, "🛠️", "Project Foundation", [
        "ASP.NET Core 10 MVC targeting net10.0",
        "EF Core 10 & Identity 10.0.7 packages",
        "WebApplicationBuilder pattern in Program.cs",
        "Empty Services, Repositories, and DTOs directories"
    ], ACCENT_INDIGO)

    feature_card(s3, 4.75, 2.3, 3.8, 3.5, "🔒", "Authentication & Cookies", [
        "ASP.NET Core Identity cookie authentication",
        "Cookie name: CipherVault.Auth (no __Host- prefix)",
        "SecurePolicy = Always (locked for dev & prod)",
        "Lockout policy: 5 failed attempts → 15 min lockout"
    ], SUCCESS)

    feature_card(s3, 8.8, 2.3, 3.8, 3.5, "📱", "UI Layout & Controllers", [
        "Responsive Bootstrap 5 layout (_Layout.cshtml)",
        "Navigation: Home, Dashboard, Vault, About",
        "HomeController for static pages",
        "DashboardController protected with [Authorize]"
    ], CYAN_GLOW)

    add_insight_box(
        s3,
        "PROGRAM.CS COMPLIANCE",
        "Avoided all deprecated APIs such as WithOpenApi, IActionContextAccessor, and WebHostBuilder, using modern ASP.NET Core 10 patterns exclusively."
    )
    add_footer(s3, 3)
    add_notes(s3, "In Phase 1, we initialized the ASP.NET Core 10 project, set up Identity cookie auth with unconditional SecurePolicy = Always, and built the responsive Bootstrap 5 layout.")

    # --------------------------------------------------------------------------
    # SLIDE 4: PHASE 2 — FORMS, VALIDATION & CRYPTO
    # --------------------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    new_slide()
    s4.background.fill.solid(); s4.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s4, "PHASE 2 • INTERACTION", "Phase 2: Forms, Data Validation & Password Generator", "Objective: Implement secure forms, DTO separation, and client-side cryptographic utilities", badge_color=ACCENT_VIOLET)

    feature_card(s4, 0.7, 2.3, 3.8, 3.5, "📝", "DTO Design Rules", [
        "Strict separation of Create and Update DTOs",
        "VaultItemCreateDTO: Secret required, Notes optional",
        "VaultItemUpdateDTO: Secret & Notes nullable",
        "Blank-Preserve Rule implemented in Service layer"
    ], ACCENT_VIOLET)

    feature_card(s4, 4.75, 2.3, 3.8, 3.5, "🛡️", "Validation & Attributes", [
        "Data annotations ([Required], [StringLength], [Url])",
        "Custom ValidCategoryAttribute for allowed categories",
        "OptionalUrlAttribute for secure URL formatting",
        "Anti-forgery tokens enabled globally on MVC forms"
    ], SUCCESS)

    feature_card(s4, 8.8, 2.3, 3.8, 3.5, "🎲", "Crypto-Secure Generator", [
        "Client-side password generator using Web Crypto API",
        "crypto.getRandomValues() (NOT Math.random())",
        "Live strength meter updating dynamically",
        "Configurable length and character sets"
    ], CYAN_GLOW)

    add_insight_box(
        s4,
        "BLANK-PRESERVE SEMANTICS",
        "The DTO design ensures that submitting an update with a blank secret preserves the existing encrypted database value without accidental overwrites."
    )
    add_footer(s4, 4)
    add_notes(s4, "Phase 2 covered forms, DTO separation (Create vs Update), validation attributes, and a cryptographically secure password generator using crypto.getRandomValues().")

    # --------------------------------------------------------------------------
    # SLIDE 5: PHASE 3 — DATABASE & ENCRYPTED CRUD
    # --------------------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    new_slide()
    s5.background.fill.solid(); s5.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s5, "PHASE 3 • PERSISTENCE", "Phase 3: Database Integration & Encrypted CRUD", "Objective: Establish EF Core persistence, repository pattern, and AES encryption at rest", badge_color=SUCCESS)

    feature_card(s5, 0.7, 2.3, 3.8, 3.5, "🗄️", "EF Core & Repositories", [
        "ApplicationDbContext inheriting IdentityDbContext",
        "DbSet<VaultItem> with User relationship (Cascade delete)",
        "Mandatory Repository pattern for VaultItem data access",
        "Migrations applied: InitialIdentitySchema & AddVaultItems"
    ], SUCCESS)

    feature_card(s5, 4.75, 2.3, 3.8, 3.5, "🔐", "Data Protection API", [
        "IDataProtectionProvider injected into VaultItemService",
        "Purpose strings: CipherVault.Secret & CipherVault.Notes",
        "Keys persisted to App_Data/keys (git-ignored)",
        "Secrets encrypted before save, decrypted on reveal only"
    ], CYAN_GLOW)

    feature_card(s5, 8.8, 2.3, 3.8, 3.5, "🚫", "IDOR & Security", [
        "Ownership verified in Service layer on every query",
        "Non-owned item access returns 404 (not 403)",
        "Prevents existence leaking across user accounts",
        "Audit logging for Create, Update, Delete, and Reveal"
    ], DANGER)

    add_insight_box(
        s5,
        "IDOR & PURPOSE STRING CONTRACT",
        "Returning 404 instead of 403 on unauthorized ID access prevents attackers from enumerating valid IDs, while immutable purpose strings protect existing ciphertexts."
    )
    add_footer(s5, 5)
    add_notes(s5, "Phase 3 introduced EF Core 10 persistence, the repository pattern, Data Protection encryption with immutable purpose strings, and IDOR-safe 404 error handling.")

    # --------------------------------------------------------------------------
    # SLIDE 6: PHASE 4 — REST API, JWT & DEPLOYMENT
    # --------------------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    new_slide()
    s6.background.fill.solid(); s6.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s6, "PHASE 4 • API & GO-LIVE", "Phase 4: REST API, JWT Authentication & Deployment", "Objective: Expose a secure REST API with JWT Bearer auth and deploy to live production hosting", badge_color=CYAN_GLOW)

    feature_card(s6, 0.7, 2.3, 3.8, 3.5, "🔌", "REST API & JWT", [
        "VaultApiController under /api/v1/vault",
        "AuthApiController handling POST /api/auth/login",
        "Dual auth scheme: Cookie (MVC) + JWT Bearer (API)",
        "API returns JSON 401 on missing auth (no redirect)"
    ], CYAN_GLOW)

    feature_card(s6, 4.75, 2.3, 3.8, 3.5, "🛡️", "Error Handling & Polish", [
        "Global ExceptionHandlingMiddleware for JSON/HTML errors",
        "Custom 404 and 500 error pages",
        "Linear UI theme with Bootstrap 5 & custom CSS",
        "Copy-to-clipboard button with 10s auto-clear"
    ], ACCENT_INDIGO)

    feature_card(s6, 8.8, 2.3, 3.8, 3.5, "🚀", "Live Production", [
        "Published via dotnet publish -c Release",
        "Hosted on MonsterASP (ciphervault.tryasp.net)",
        "Let's Encrypt SSL certificate bound via HTTPS",
        "Data Protection key directory secured outside web root"
    ], SUCCESS)

    add_insight_box(
        s6,
        "DUAL AUTH SCHEME SEPARATION",
        "Configuring explicit authentication schemes (`AuthenticationSchemes = JwtBearerDefaults.AuthenticationScheme`) ensures API clients receive clean JSON 401s without HTML redirect loops."
    )
    add_footer(s6, 6)
    add_notes(s6, "Phase 4 delivered the REST API with JWT bearer auth, dual scheme separation, global error handling, and live HTTPS deployment on MonsterASP.")

    # --------------------------------------------------------------------------
    # SLIDE 7 — LIVE APP DEMO (DASHBOARD & VAULT)
    # --------------------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    new_slide()
    s7.background.fill.solid(); s7.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s7, "PROTOTYPE", "Live App Demo: Dashboard & Vault Management", "Visual verification of production UI, analytics, and server-side search", badge_color=SUCCESS)

    add_plot_frame(s7, "dashboard/05-dashboard-populated.png", "Dashboard: KPI Stats, Chart.js Breakdown & Recent Items", 0.7, 2.3, 5.7, 4.3, border_color=SUCCESS)
    add_plot_frame(s7, "vault/07-list-with-filter.png", "Vault: Server-Side Search, Category Filter & Pagination", 6.9, 2.3, 5.7, 4.3, border_color=CYAN_GLOW)

    add_notes(s7, "Demonstrating the live application UI: the dashboard with Chart.js analytics and the vault list featuring server-side search and category filtering.")
    add_footer(s7, 7)

    # --------------------------------------------------------------------------
    # SLIDE 8 — LIVE APP DEMO (ENCRYPTED CRUD & REVEAL)
    # --------------------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    new_slide()
    s8.background.fill.solid(); s8.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s8, "PROTOTYPE", "Live App Demo: Encrypted CRUD & Reveal Flow", "Secure creation, masking by default, and authorized on-demand decryption", badge_color=SUCCESS)

    add_plot_frame(s8, "vault/09-create-form.png", "Create Form with Password Generator & Strength Meter", 0.7, 2.3, 5.7, 4.3, border_color=ACCENT_VIOLET)
    add_plot_frame(s8, "vault/12-details-revealed.png", "Details View: Masked Secrets & On-Demand Reveal", 6.9, 2.3, 5.7, 4.3, border_color=SUCCESS)

    add_notes(s8, "Showing the create form with the crypto-secure generator and the details view where secrets remain masked by default until explicitly revealed by the owner.")
    add_footer(s8, 8)

    # --------------------------------------------------------------------------
    # SLIDE 9 — SECURITY & CRYPTO VERIFICATION
    # --------------------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    new_slide()
    s9.background.fill.solid(); s9.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s9, "SECURITY PROOF", "Security Verification & Database Inspection", "SSMS verification proving ciphertext-only storage and cross-user isolation", badge_color=DANGER)

    add_plot_frame(s9, "security/15-db-encrypted-ssms.png", "SSMS Proof: EncryptedSecret & EncryptedNotes Store Ciphertext Only", 0.7, 2.3, 5.7, 4.3, border_color=DANGER)
    add_plot_frame(s9, "security/18-devtools-hsts.png", "DevTools Proof: HSTS Header & Secure Cookies Active", 6.9, 2.3, 5.7, 4.3, border_color=SUCCESS)

    add_notes(s9, "Database inspection via SSMS confirms that plaintext secrets never touch the database. Browser DevTools confirm secure cookies and HSTS headers in production.")
    add_footer(s9, 9)

    # --------------------------------------------------------------------------
    # SLIDE 10 — REST API & POSTMAN VERIFICATION
    # --------------------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    new_slide()
    s10.background.fill.solid(); s10.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s10, "API VERIFICATION", "REST API Verification via Postman", "End-to-end testing of JWT login, unauthorized responses, and IDOR protection", badge_color=CYAN_GLOW)

    add_plot_frame(s10, "api/20-postman-login.png", "Postman: POST /api/auth/login Issuing JWT Bearer Token", 0.7, 2.3, 3.7, 4.3, border_color=CYAN_GLOW)
    add_plot_frame(s10, "api/22-postman-401.png", "Postman: HTTP 401 Unauthorized (No Redirect)", 4.8, 2.3, 3.7, 4.3, border_color=WARNING)
    add_plot_frame(s10, "api/23-postman-404-idor.png", "Postman: HTTP 404 Not Found (IDOR Cross-User Guard)", 8.9, 2.3, 3.7, 4.3, border_color=DANGER)

    add_notes(s10, "Postman test results confirming token issuance, JSON 401 on missing tokens, and 404 Not Found on cross-user IDOR access attempts.")
    add_footer(s10, 10)

    # --------------------------------------------------------------------------
    # SLIDE 11 — COMPREHENSIVE TESTING MATRIX
    # --------------------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    new_slide()
    s11.background.fill.solid(); s11.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s11, "VERIFICATION", "Comprehensive Testing & Verification Matrix", "Summary of 11 critical test cases verified across local development and production", badge_color=SUCCESS)

    t11_shape = s11.shapes.add_table(12, 3, Inches(0.7), Inches(2.3), Inches(11.933), Inches(4.3))
    t11 = t11_shape.table
    headers11 = ["Test Case Description", "Expected Behavior / Requirement", "Status"]
    rows11 = [
        ["IDOR Cross-User Access", "Attempting to access another user's vault item returns 404 Not Found", "PASS ✔"],
        ["Blank Secret on Update", "Submitting an edit form with a blank secret preserves existing ciphertext", "PASS ✔"],
        ["Blank Notes on Update", "Leaving notes blank preserves them; empty string clears them to null", "PASS ✔"],
        ["API Missing Token", "Calling /api/v1/vault without a token returns HTTP 401 JSON (no redirect)", "PASS ✔"],
        ["API ?decrypt=true", "Supplying ?decrypt=true query parameter is rejected with HTTP 400", "PASS ✔"],
        ["Account Lockout", "5 failed login attempts trigger a 15-minute account lockout (HTTP 423)", "PASS ✔"],
        ["HSTS Header", "Production environment emits Strict-Transport-Security header", "PASS ✔"],
        ["Encryption at Rest", "Database inspection confirms only ciphertext is stored in Secret/Notes columns", "PASS ✔"],
        ["Key Persistence", "Application restart does not invalidate existing ciphertext (Key ring persists)", "PASS ✔"],
        ["Global Exception Handling", "Unhandled exceptions return JSON 500 for API and custom error page for MVC", "PASS ✔"],
        ["Mobile Responsiveness", "Bootstrap 5 layout adapts cleanly across mobile and desktop viewports", "PASS ✔"]
    ]
    for c_i, h in enumerate(headers11):
        t11.cell(0, c_i).text = h
    for r_i, r_data in enumerate(rows11, start=1):
        for c_i, val in enumerate(r_data):
            t11.cell(r_i, c_i).text = val
    style_table(t11, [3.2, 7.333, 1.4], font_size=9.5)
    for r_i in range(1, 12):
        t11.cell(r_i, 2).text_frame.paragraphs[0].font.color.rgb = SUCCESS
        t11.cell(r_i, 2).text_frame.paragraphs[0].font.bold = True
        t11.cell(r_i, 2).text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

    add_notes(s11, "Summary table of all 11 critical test cases, confirming 100% pass rate across security, CRUD, API, and deployment verifications.")
    add_footer(s11, 11)

    # --------------------------------------------------------------------------
    # SLIDE 12: TECHNICAL CHALLENGES & SOLUTIONS
    # --------------------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    new_slide()
    s12.background.fill.solid(); s12.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s12, "RETROSPECTIVE", "Technical Challenges & Solutions", "Real-world hurdles encountered during development and how they were resolved", badge_color=WARNING)

    challenges = [
        ("__Host- Cookie Prefix Conflict", "The __Host- cookie prefix caused login failures during local HTTP development. Resolved by renaming the cookie to CipherVault.Auth while maintaining unconditional SecurePolicy = Always."),
        ("Lockout Parity Across Interfaces", "MVC login and API login had divergent lockout behaviors. Resolved by scaffolding the Identity Login page and enabling lockoutOnFailure: true across both pathways."),
        ("HSTS Poisoning on Localhost", "Enabling HSTS globally caused browser HSTS pinning on local development ports. Resolved by gating UseHsts() strictly inside if (app.Environment.IsProduction())."),
        ("Key Ring Persistence on Shared Host", "Default key directories risked being overwritten during deployments. Resolved by enforcing DataProtection__KeyDirectory outside the web root via environment variables."),
        ("DTO Blank-Preserve Semantics", "Distinguishing between 'keep existing secret' and 'clear field' in updates. Resolved by implementing explicit null vs empty string checks in the VaultItemService.")
    ]
    y = 2.3
    for title, desc in challenges:
        add_card(s12, 0.7, y, 11.933, 0.82, border_color=BORDER_SLATE, fill_color=CARD_BG)
        add_card(s12, 0.7, y, 0.08, 0.82, fill_color=WARNING)
        add_text(s12, title, 1.0, y + 0.1, 3.5, 0.35, size=11.5, bold=True, color=TEXT_WHITE)
        add_text(s12, desc, 4.6, y + 0.1, 7.8, 0.65, size=10.5, color=TEXT_MUTED)
        y += 0.94

    add_notes(s12, "Discussing technical challenges: cookie prefix issues, lockout parity, HSTS gating, key persistence on shared hosting, and DTO blank-preserve semantics.")
    add_footer(s12, 12)

    # --------------------------------------------------------------------------
    # SLIDE 13: LESSONS LEARNED & BEST PRACTICES
    # --------------------------------------------------------------------------
    s13 = prs.slides.add_slide(blank_layout)
    new_slide()
    s13.background.fill.solid(); s13.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s13, "REFLECTION", "Lessons Learned & Engineering Best Practices", "Key takeaways from building a production-oriented ASP.NET Core 10 application", badge_color=ACCENT_VIOLET)

    lessons = [
        ("Architecture Discipline", "Enforcing strict separation between controllers, services, and repositories prevents spaghetti code and simplifies unit testing."),
        ("Cryptographic Contracts", "Purpose strings are immutable contracts. Planning key management and directory persistence before production deployment is mandatory."),
        ("Security by Default", "IDOR protection belongs in the service layer, not controller actions, ensuring ownership checks cannot be accidentally bypassed."),
        ("Dual Auth Complexity", "Mixing cookie authentication and JWT bearer tokens requires explicit scheme configuration to prevent unwanted login redirects on API endpoints."),
        ("Environment Parity", "Testing production settings (HSTS, secure cookies, external key directories) locally via user-secrets and environment profiles prevents release-day surprises.")
    ]
    y = 2.3
    for i, (title, desc) in enumerate(lessons, 1):
        add_card(s13, 0.7, y, 11.933, 0.85, border_color=BORDER_SLATE, fill_color=CARD_BG)
        circ = s13.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(y + 0.22), Inches(0.4), Inches(0.4))
        circ.fill.solid(); circ.fill.fore_color.rgb = ACCENT_VIOLET
        circ.line.fill.background()
        circ.text_frame.text = str(i)
        p = circ.text_frame.paragraphs[0]
        p.font.size = Pt(11); p.font.bold = True; p.font.color.rgb = BG_OBSIDIAN
        p.alignment = PP_ALIGN.CENTER
        circ.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

        add_text(s13, title, 1.5, y + 0.12, 3.5, 0.35, size=12, bold=True, color=TEXT_WHITE)
        add_text(s13, desc, 5.1, y + 0.12, 7.3, 0.65, size=11, color=TEXT_MUTED)
        y += 0.96

    add_notes(s13, "Summarizing lessons learned: architecture discipline, immutable cryptographic purpose strings, service-layer IDOR enforcement, and dual auth complexity.")
    add_footer(s13, 13)

    # --------------------------------------------------------------------------
    # SLIDE 14: FUTURE ENHANCEMENTS
    # --------------------------------------------------------------------------
    s14 = prs.slides.add_slide(blank_layout)
    new_slide()
    s14.background.fill.solid(); s14.background.fill.fore_color.rgb = BG_OBSIDIAN
    add_header(s14, "ROADMAP", "Future Enhancements & Production Roadmap", "Pathways toward enterprise-grade security and operational maturity", badge_color=CYAN_GLOW)

    feature_card(s14, 0.7, 2.3, 3.8, 4.3, "🛡️", "Security Enhancements", [
        "Multi-factor authentication (TOTP / WebAuthn)",
        "Secure password reset and recovery flows",
        "Azure Key Vault / AWS Secrets Manager integration",
        "Independent third-party security audit"
    ], CYAN_GLOW)

    feature_card(s14, 4.75, 2.3, 3.8, 4.3, "🧪", "Quality & Interoperability", [
        "Automated unit & integration test suite (xUnit)",
        "Encrypted vault import / export (JSON / CSV)",
        "Self-hosted Chart.js assets for offline operation",
        "Automated SBOM and dependency vulnerability scanning"
    ], ACCENT_INDIGO)

    feature_card(s14, 8.8, 2.3, 3.8, 4.3, "⚙️", "Operations & DevOps", [
        "Paid hosting migration (no app sleep, auto SSL)",
        "CI/CD deployment pipelines via GitHub Actions",
        "Automated database and key ring backup jobs",
        "Centralized structured logging (Serilog + Seq)"
    ], SUCCESS)

    add_notes(s14, "Outlining the future roadmap: MFA, automated testing, vault import/export, CI/CD pipelines, and cloud key vault integration.")
    add_footer(s14, 14)

    # --------------------------------------------------------------------------
    # SLIDE 15: CONCLUSION & Q&A
    # --------------------------------------------------------------------------
    s15 = prs.slides.add_slide(blank_layout)
    new_slide()
    s15.background.fill.solid(); s15.background.fill.fore_color.rgb = BG_OBSIDIAN

    add_card(s15, 0.55, 0.55, 12.233, 6.2, border_color=SUCCESS, fill_color=CARD_BG)
    add_card(s15, 0.9, 0.88, 6.5, 0.42, border_color=CYAN_GLOW, fill_color=CARD_ALT)
    add_text(s15, f"{COURSE_CODE}  •  ALL 4 PHASES SUCCESSFULLY COMPLETED", 0.95, 0.93, 6.4, 0.32, size=11, color=CYAN_GLOW, bold=True, align=PP_ALIGN.CENTER)

    add_text(s15, "Thank You! Questions & Discussion", 0.9, 1.48, 11.5, 0.7, size=32, color=TEXT_WHITE, bold=True)
    add_text(s15, f"Live Production Build: {LIVE_URL}   |   Source Code: {GITHUB_URL}", 0.9, 2.22, 11.5, 0.4, size=13.5, color=SUCCESS, bold=True)

    badges = [
        ("PHASE 1 • MVC & UI", "ASP.NET Core 10 MVC\nBootstrap 5 Responsive", ACCENT_INDIGO),
        ("PHASE 2 • FORMS", "DTO Separation\nCrypto-Secure Generator", ACCENT_VIOLET),
        ("PHASE 3 • CRUD & DB", "EF Core 10 + AES Data Protection\nIDOR-Safe 404 Guard", SUCCESS),
        ("PHASE 4 • API & LIVE", "REST API + JWT Bearer\nDeployed on MonsterASP", CYAN_GLOW)
    ]
    for i, (b_top, b_sub, col) in enumerate(badges):
        bx = 0.9 + i * 2.92
        add_card(s15, bx, 2.9, 2.72, 1.42, border_color=col, fill_color=CARD_ALT)
        add_text(s15, b_top, bx + 0.08, 3.0, 2.56, 0.35, size=12.5, color=col, bold=True, align=PP_ALIGN.CENTER)
        add_text(s15, b_sub, bx + 0.08, 3.4, 2.56, 0.8, size=10, color=TEXT_WHITE, align=PP_ALIGN.CENTER)

    add_card(s15, 0.9, 4.52, 11.55, 1.95, border_color=BORDER_SLATE, fill_color=CARD_ALT)
    add_text(
        s15,
        f"✔ Phase 1, 2, 3 & 4 Implementation Reports Complete     |     ✔ Live HTTPS Deployment Verified     |     ✔ Postman Collection & Docs Ready\n\n"
        f"Student Name: {STUDENT_NAME} (Roll No: {ROLL_NO})     •     Project Guide: {GUIDE_NAME}\n"
        f"Course Title: {COURSE_NAME} ({COURSE_CODE})     •     Institution: {COLLEGE}",
        1.05, 4.75, 11.2, 1.5, size=11.5, color=ACCENT_VIOLET, bold=True, align=PP_ALIGN.CENTER
    )

    add_footer(s15, 15)
    add_notes(s15, "Thank you Sir/Ma'am. I am now ready for your questions and a live demonstration of CipherVault.")

    prs.save(OUTPUT_FILE)
    print("=" * 75)
    print(f"✅ PROFESSIONAL CIPHERVAULT PRESENTATION GENERATED: {os.path.abspath(OUTPUT_FILE)}")
    print(f"   • Total Slides: {slide_num}")
    print(f"   • Guide: {GUIDE_NAME} | Course Code: {COURSE_CODE}")
    print(f"   • Theme: Linear Dark Vault (App-Aligned & Professional)")
    print("=" * 75)

if __name__ == "__main__":
    build_presentation()